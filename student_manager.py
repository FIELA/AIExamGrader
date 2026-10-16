# Copyright (c) 2025 JASim. Licensed under NCEL-Strict License v2.0.
# STRICT NON-COMMERCIAL USE ONLY. No AI/ML training, fine-tuning, or public distribution of Derivative Works.
# Modifications may only be shared as Patch Files.
# Public forks allowed solely for PRs (delete within 14 days after PR merged, rejected, or closed).
# Commercial licensing inquiries: nicofiela@outlook.com. See LICENSE file for full terms.

import os
import re
import csv
import unicodedata
try:
    import openpyxl
except ImportError:
    openpyxl = None
from typing import List, Dict, Optional, Tuple

# A file named exactly "Room-Seat" (e.g. 01-05.jpg); dated scanner names like 2026-10-16-001 do not qualify
ROOM_SEAT_NAME = re.compile(r"^(\d+)\s*-\s*(\d+)$")

# What OCR returns for a field that was left blank
EMPTY_TEXT = {'', 'none', 'null', 'nan', 'n/a', 'na', '无', '空', '未知', '未填', '未填写'}
CN_DIGITS = {'零': 0, '〇': 0, '一': 1, '二': 2, '两': 2, '三': 3, '四': 4, '五': 5, '六': 6, '七': 7, '八': 8, '九': 9}

# Matches that a teacher should look at before trusting the student shown
REVIEW_METHODS = {'class_name_similar', 'ambiguous', 'unmatched'}

def cell_text(value) -> str:
    """Spreadsheet cell as text: empty cells become '' and whole floats such as 6.0 become '6'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()

def same_number(a, b) -> bool:
    """Compare room/seat values so that '06' equals '6'."""
    a, b = str(a).strip(), str(b).strip()
    if a.isdecimal() and b.isdecimal():
        return int(a) == int(b)
    return a == b

def clean_text(value) -> str:
    """OCR or roster text without spaces; placeholders such as 'None' or '无' become ''."""
    text = re.sub(r"\s+", "", unicodedata.normalize('NFKC', str(value if value is not None else '')))
    return '' if text.lower() in EMPTY_TEXT else text

def name_key(value) -> str:
    """Name in comparable form: letters plus the middle dot of names such as 阿依古丽·买买提."""
    text = clean_text(value).replace('•', '·').replace('・', '·')
    return re.sub(r"[^\w·]|[\d_]", "", text).lower()

def _number(text: str) -> int:
    """'七' -> 7, '十一' -> 11, '二十' -> 20, '七〇一' -> 701, '3' -> 3."""
    if text.isdecimal():
        return int(text)
    if '十' in text:
        tens, _, ones = text.partition('十')
        return (CN_DIGITS.get(tens, 1) if tens else 1) * 10 + (CN_DIGITS.get(ones, 0) if ones else 0)
    return int(''.join(str(CN_DIGITS[c]) for c in text))

def class_parts(value) -> Optional[Tuple[Optional[int], int]]:
    """(grade, class number) of a label such as 701, 7(1)班, 七年级一班 or 初一1班; grade is None for 3班."""
    text = clean_text(value)
    # 初一 is grade 7 and 高一 grade 10
    text = re.sub(r"([初高])中?([一二三1-3])", lambda m: f"{(6 if m.group(1) == '初' else 9) + _number(m.group(2))}年级", text)
    text = re.sub(r"[零〇一二两三四五六七八九十]+", lambda m: str(_number(m.group())), text)
    numbers = re.findall(r"\d+", text)
    if not numbers:
        return None
    if len(numbers) == 1:
        n = numbers[0]
        # 701 and 0701 are grade 7 class 1
        return (int(n[:-2]), int(n[-2:])) if len(n) >= 3 else (None, int(n))
    return (int(numbers[-2]), int(numbers[-1]))

def same_class(a, b) -> bool:
    """Whether two class labels name the same class, e.g. 701, 7(1)班 and 七年级一班."""
    pa, pb = class_parts(a), class_parts(b)
    if pa is None or pb is None:
        ta, tb = clean_text(a).rstrip('班'), clean_text(b).rstrip('班')
        return bool(ta) and ta == tb
    return pa[1] == pb[1] and (pa[0] is None or pb[0] is None or pa[0] == pb[0])

def edit_distance(a: str, b: str) -> int:
    """Number of single-character insertions, deletions or substitutions between two strings."""
    previous = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        current = [i]
        for j, cb in enumerate(b, 1):
            current.append(min(previous[j] + 1, current[j - 1] + 1, previous[j - 1] + (ca != cb)))
        previous = current
    return previous[-1]

def student_key(student: Dict) -> Tuple[str, ...]:
    """Identity of a roster entry."""
    return tuple(str(student.get(k, '')) for k in ('name', 'class', 'id', 'room', 'seat'))

class StudentManager:
    def __init__(self):
        self.students: List[Dict[str, str]] = []
        self.student_path: str = ""

    def load_from_file(self, path: str) -> int:
        """
        Load students from an Excel or CSV file.
        Returns the number of students loaded.
        """
        self.students = []
        self.student_path = path
        rows = []
        
        try:
            if path.endswith('.csv'):
                with open(path, 'r', encoding='utf-8-sig') as f:
                    reader = csv.reader(f)
                    rows = list(reader)
            else:
                if openpyxl is None:
                    raise ImportError("openpyxl is not installed. Please install it to read Excel files.")
                wb = openpyxl.load_workbook(path, data_only=True)
                ws = wb.active
                for row in ws.iter_rows(values_only=True):
                    rows.append(list(row))
            
            if not rows:
                return 0

            header = [cell_text(h) for h in rows[0]]
            idx_map = {}
            for i, col_name in enumerate(header):
                col_lower = col_name.lower()
                if "考号" in col_name or "学号" in col_name or "准考证号" in col_name or "id" in col_lower: idx_map['id'] = i
                elif "姓名" in col_name or "name" in col_lower: idx_map['name'] = i
                elif "班级" in col_name or "class" in col_lower: idx_map['class'] = i
                elif "考场" in col_name or "room" in col_lower: idx_map['room'] = i
                elif "座号" in col_name or "座位" in col_name or "seat" in col_lower: idx_map['seat'] = i

            for row in rows[1:]:
                values = [cell_text(v) for v in row]
                if not any(values):
                    continue  # Blank line
                s = {'id': '', 'name': '', 'class': '', 'room': '', 'seat': ''}
                for key, i in idx_map.items():
                    if i < len(values): s[key] = values[i]

                if s['room'] and s['room'].isdigit(): s['room'] = s['room'].zfill(2)
                
                self.students.append(s)
                
            return len(self.students)
        except Exception as e:
            raise Exception(f"Failed to parse student file: {str(e)}")

    @staticmethod
    def parse_room_seat(filename: str) -> Tuple[str, str, bool]:
        """
        Parse a "Room-Seat" filename such as 01-05.jpg or 01-05缺.jpg.
        Returns (room, seat, is_absent); room and seat are '' for any other name.
        """
        name_part = os.path.splitext(os.path.basename(filename))[0]
        is_absent = "缺" in name_part
        match = ROOM_SEAT_NAME.match(name_part.replace("缺", "").strip(" -_"))
        if not match:
            return "", "", is_absent
        return match.group(1), match.group(2), is_absent

    def get_student_by_filename(self, filename: str) -> Tuple[Dict[str, str], bool]:
        """
        Parse filename and find student.
        Returns (student_dict, is_absent)
        """
        room_no, seat_no, is_absent = self.parse_room_seat(filename)

        found_student = None
        if self.students and room_no and seat_no:
            for s in self.students:
                if same_number(s['room'], room_no) and same_number(s['seat'], seat_no):
                    found_student = dict(s)  # Callers annotate it; keep the roster entry clean
                    break
        
        if not found_student:
            found_student = {
                'id': '未知', 'name': '未知', 'class': '未知',
                'room': room_no, 'seat': seat_no
            }

        return found_student, is_absent

    def resolve_student(self, filename: str, ocr: Optional[Dict] = None) -> Tuple[Dict[str, str], Dict]:
        """
        Find the roster entry for a graded answer sheet. Tries, in order: a Room-Seat file name,
        the filled then the written exam number, the written Room-Seat, class + name, the name alone
        when only one student has it, and a name one character away within the written class.
        Returns (student, match): a copy of the roster entry (or the '未知' placeholder) and
        {'method': ..., 'candidates': [...]}, where candidates are the students to choose from in review.
        """
        ocr = ocr or {}
        student, _ = self.get_student_by_filename(filename)
        if student['name'] != '未知':
            return student, {'method': 'filename', 'candidates': []}

        def matched(entry, method):
            return dict(entry), {'method': method, 'candidates': []}

        def unresolved(method, entries):
            return student, {'method': method, 'candidates': [dict(s) for s in entries[:3]]}

        for field, method in (('ocr_id_filled', 'id_filled'), ('ocr_id_written', 'id_written')):
            exam_id = clean_text(ocr.get(field))
            if exam_id.isdecimal():
                found = [s for s in self.students if s['id'] and same_number(s['id'], exam_id)]
                if len(found) == 1:
                    return matched(found[0], method)

        room, seat = clean_text(ocr.get('ocr_room')), clean_text(ocr.get('ocr_seat'))
        if room.isdecimal() and seat.isdecimal():
            found = [s for s in self.students if s['room'] and s['seat'] and same_number(s['room'], room) and same_number(s['seat'], seat)]
            if len(found) == 1:
                return matched(found[0], 'room_seat')

        name = name_key(ocr.get('ocr_name'))
        if not name:
            return unresolved('unmatched', [])
        klass = clean_text(ocr.get('ocr_class'))
        in_class = [s for s in self.students if klass and same_class(s['class'], klass)]

        exact = [s for s in in_class if name_key(s['name']) == name]
        if len(exact) == 1:
            return matched(exact[0], 'class_name')
        if exact:
            return unresolved('ambiguous', exact)  # Two students of one class share the name

        # A misread character within the class competes with the same name in another class
        similar = self._similar(name, in_class)
        elsewhere = [s for s in self.students if name_key(s['name']) == name]
        if len(elsewhere) == 1 and not similar:
            return matched(elsewhere[0], 'name')
        if len(similar) == 1 and not elsewhere:
            return matched(similar[0], 'class_name_similar')
        if similar or elsewhere:
            return unresolved('ambiguous', similar + elsewhere)
        return unresolved('unmatched', self._similar(name, self.students))

    @staticmethod
    def _similar(name: str, pool: List[Dict[str, str]]) -> List[Dict[str, str]]:
        """Roster entries whose name differs from name by one character (names of two or more characters)."""
        if len(name) < 2:
            return []
        return [s for s in pool if len(name_key(s['name'])) >= 2 and edit_distance(name_key(s['name']), name) == 1]

    def match_summary(self, records):
        """
        records: (image name, bound student, match method) for each graded sheet.
        Returns (images to check, {student key: images} for students on several sheets,
        roster students that no sheet is bound to).
        """
        to_check, by_student = [], {}
        for image, student, method in records:
            if student.get('name') in (None, '', '未知'):
                to_check.append(image)
                continue
            if method in REVIEW_METHODS:
                to_check.append(image)
            by_student.setdefault(student_key(student), []).append(image)
        several = {key: images for key, images in by_student.items() if len(images) > 1}
        without_sheet = [s for s in self.students if student_key(s) not in by_student]
        return to_check, several, without_sheet
