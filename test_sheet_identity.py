"""Answer-sheet file names, roster matching and report ownership, run headless with a fake grader."""
import json
import os
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl
from PIL import Image

from AutoGrader import App
from report_files import find_report, unique_report_prefix
from review_window import ReviewWindow
from student_manager import StudentManager

ROSTER = [
    {'id': '23090828', 'name': '张三', 'class': '701', 'room': '01', 'seat': '5'},
    {'id': '23090829', 'name': '李四', 'class': '701', 'room': '01', 'seat': '6'},
]


class SheetApp:
    """Runs the App's real file-handling methods without opening a window."""
    t = App.t
    get_folder_path = App.get_folder_path
    _image_folders = App._image_folders
    _rename_new_arrivals = App._rename_new_arrivals
    process_single_file = App.process_single_file
    save_markdown = App.save_markdown
    generate_report_content = App.generate_report_content
    match_label = App.match_label
    is_absent_record = staticmethod(App.is_absent_record)
    regrade_single_file = App.regrade_single_file
    on_review_save = App.on_review_save
    log_student_matches = App.log_student_matches
    rematch_students = App.rematch_students
    ensure_jsons_exist = App.ensure_jsons_exist
    open_review_window = App.open_review_window
    current_lang = "CN"
    layout_description = None
    rubric_path = "rubric.txt"

    def __init__(self, exam_folder, students=None):
        self.exam_folder = str(exam_folder)
        self.student_manager = StudentManager()
        self.student_manager.students = [dict(s) for s in students or []]
        self.stop_event = threading.Event()
        self.write_lock = threading.RLock()
        self.rename_lock = threading.Lock()
        self.completed_count = self.session_completed_count = self.total_files = 0
        self.logs, self.csv_rows = [], []

    def ui_after(self, delay, callback, *args): callback(*args)
    def log(self, message): self.logs.append(message)
    def write_summary_csv(self, data_dict, target_path=None): self.csv_rows.append(data_dict)
    def regenerate_csv_from_jsons(self): pass
    def update_progress_ui(self): pass
    def read_rubric_text(self): return "rubric"


class FakeGrader:
    """Returns the OCR fields given for each image name, like the model would."""
    def __init__(self, ocr, barrier=None):
        self.ocr, self.barrier = ocr, barrier

    def grade_exam(self, rubric_text, image_path, layout_description=None):
        fields = self.ocr[os.path.basename(image_path)]
        if self.barrier:
            self.barrier.wait(timeout=5)  # Let parallel workers reach the rename step together
        details = [{"question_id": "1", "type": "客观题", "student_answer": "A", "score": 8, "max_score": 8}]
        return dict(fields, total_score=8, details=details)


def ocr(name="", klass="701", exam_id=""):
    return {"ocr_name": name, "ocr_class": klass, "ocr_id_written": exam_id, "ocr_id_filled": exam_id}


class ReviewStub:
    """Runs the review window's real binding and save methods without opening a window."""
    student_display = staticmethod(ReviewWindow.student_display)
    bind_selected_student = ReviewWindow.bind_selected_student
    save_to_disk = ReviewWindow.save_to_disk
    update_markdown_file = ReviewWindow.update_markdown_file

    def __init__(self, app, json_path):
        self.parent_app = app
        self.student_manager = app.student_manager
        self.reports_dir = app.get_folder_path('reports')
        self.current_json_path = str(json_path)
        self.current_data = json.loads(Path(json_path).read_text(encoding='utf-8'))
        self.report_index = {}
        self.image_files, self.current_index = [], 0
        self.chk_absence_var = type("Var", (), {"get": lambda self: False})()

    def update_score_display(self): pass
    def update_csv_row(self): pass


class TestSheetIdentity(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.exam = Path(self.tmp.name) / "exam"
        self.exam.mkdir()
        self.reports = self.exam / "阅卷报告"
        self.success = self.exam / "成功归档"

    def tearDown(self):
        self.tmp.cleanup()

    def write_image(self, path, shade):
        Image.new("RGB", (8, 8), (shade, 0, 0)).save(path)
        return Path(path).read_bytes()

    def grade(self, sheets, students=None, workers=1, barrier=None):
        """Grade every image in sheets (name -> OCR fields); return the app and each original image's bytes."""
        scans = {name: self.write_image(self.exam / name, 30 * i) for i, name in enumerate(sheets)}
        app = SheetApp(self.exam, students)
        grader = FakeGrader(sheets, barrier)
        with ThreadPoolExecutor(max_workers=workers) as pool:
            done = list(pool.map(lambda name: app.process_single_file(grader, "rubric", name, str(self.exam), False, 1), sheets))
        self.assertTrue(all(done), app.logs)
        return app, scans

    def report_owners(self):
        """Report file -> (original image, renamed image)."""
        owners = {}
        for f in sorted(os.listdir(self.reports)):
            if f.endswith('.json'):
                data = json.loads((self.reports / f).read_text(encoding='utf-8'))
                owners[f] = (data['original_filename'], data.get('renamed_filename'))
        return owners

    def assert_each_sheet_kept(self, scans):
        """Every scan has exactly one report, and that report's image in success is that scan."""
        owners = self.report_owners()
        self.assertEqual(sorted(original for original, _ in owners.values()), sorted(scans))
        for report, (original, renamed) in owners.items():
            image = renamed or original
            self.assertEqual(report, os.path.splitext(image)[0] + ".json")
            self.assertEqual((self.success / image).read_bytes(), scans[original], image)
        self.assertEqual(len(os.listdir(self.success)), len(scans))

    # A: scanner names, no exam number on the sheet
    def test_scans_without_exam_numbers_keep_their_names(self):
        app, scans = self.grade({"IMG_0001.jpg": ocr("张三"), "IMG_0002.jpg": dict(ocr("李四"), ocr_id_filled=None)})
        self.assertEqual(self.report_owners(), {"IMG_0001.json": ("IMG_0001.jpg", None), "IMG_0002.json": ("IMG_0002.jpg", None)})
        self.assert_each_sheet_kept(scans)

    # B: dashes in scanner names are not Room-Seat
    def test_dated_scan_names_are_not_read_as_room_seat(self):
        app, scans = self.grade({"2026-10-16-001.jpg": ocr("王五", "702"), "2026-10-16-002.jpg": ocr("张三", "702")})
        self.assert_each_sheet_kept(scans)
        self.assertEqual([(row['考场'], row['座号']) for row in app.csv_rows], [('', ''), ('', '')])

    # C: the model writes a placeholder instead of leaving the exam number empty
    def test_placeholder_exam_numbers_do_not_share_a_report(self):
        app, scans = self.grade({"IMG_0003.jpg": ocr("张三", exam_id="无"), "IMG_0004.jpg": ocr("李四", exam_id="无")})
        self.assertEqual(set(self.report_owners()), {"IMG_0003.json", "IMG_0004.json"})
        self.assert_each_sheet_kept(scans)

    # D: two sheets read as the same exam number, one after the other
    def test_repeated_exam_number_does_not_overwrite_the_first_sheet(self):
        app, scans = self.grade({"IMG_0005.jpg": ocr("张三", exam_id="701"), "IMG_0006.jpg": ocr("李四", exam_id="701")})
        self.assertEqual(self.report_owners(), {"701.json": ("IMG_0005.jpg", "701.jpg"), "701_1.json": ("IMG_0006.jpg", "701_1.jpg")})
        self.assert_each_sheet_kept(scans)

    # D: the same, with both workers renaming at the same moment
    def test_parallel_workers_pick_different_names(self):
        sheets = {"IMG_0005.jpg": ocr("张三", exam_id="701"), "IMG_0006.jpg": ocr("李四", exam_id="701")}
        app, scans = self.grade(sheets, workers=2, barrier=threading.Barrier(2))
        self.assertEqual(sorted(renamed for _, renamed in self.report_owners().values()), ["701.jpg", "701_1.jpg"])
        self.assert_each_sheet_kept(scans)

    # G: sheets with exam numbers are renamed, matched by that number and remember their scan; regrading keeps that record
    def test_exam_number_rename_keeps_original_scan_on_regrade(self):
        app, scans = self.grade({"IMG_0007.jpg": ocr("张三", exam_id="23090828")}, students=ROSTER)
        self.assertEqual(self.report_owners(), {"01-5.json": ("IMG_0007.jpg", "23090828.jpg")})
        app.grader_engine = FakeGrader({"23090828.jpg": ocr("张三", exam_id="23090828")})
        self.assertEqual(app.regrade_single_file(str(self.success / "23090828.jpg")), (True, "Success"))
        self.assertEqual(self.report_owners(), {"01-5.json": ("IMG_0007.jpg", "23090828.jpg")})

    # H, I: zero-padded or not, Room-Seat names find the roster entry
    def test_room_seat_names_match_roster_with_or_without_zero_padding(self):
        manager = StudentManager()
        manager.students = [dict(s) for s in ROSTER]
        for name in ("01-06.jpg", "1-6.jpg", "01-06缺.jpg", "01 - 06.jpg"):
            with self.subTest(name=name):
                student, absent = manager.get_student_by_filename(name)
                self.assertEqual(student['name'], '李四')
                self.assertEqual(absent, '缺' in name)
        for name in ("2026-10-16-001.jpg", "IMG_0001.jpg", "1-6-2.jpg"):
            with self.subTest(name=name):
                student, _ = manager.get_student_by_filename(name)
                self.assertEqual((student['name'], student['room'], student['seat']), ('未知', '', ''))
        student, _ = manager.get_student_by_filename("1-6.jpg")
        student['is_absent'] = True
        self.assertNotIn('is_absent', manager.students[1])

    def test_two_scans_of_one_seat_keep_separate_reports(self):
        app, scans = self.grade({"1-6.jpg": ocr("李四"), "01-06.jpg": ocr("李四")}, students=ROSTER)
        self.assertEqual(self.report_owners(), {"01-6.json": ("1-6.jpg", None), "01-06.json": ("01-06.jpg", None)})
        self.assertTrue(any("已属于另一张答题卡" in line for line in app.logs), app.logs)

    def test_new_scan_reusing_a_processed_name_is_graded_separately(self):
        self.success.mkdir()
        (self.exam / "原始文件").mkdir()
        first = self.write_image(self.success / "IMG_0001.jpg", 200)
        self.write_image(self.exam / "原始文件" / "IMG_0001.jpg", 200)
        self.write_image(self.exam / "原始文件" / "IMG_0002.jpg", 100)
        second = self.write_image(self.exam / "IMG_0001.jpg", 50)  # Next batch; the scanner restarted numbering
        self.write_image(self.exam / "IMG_0002.jpg", 100)  # Same sheet as its backup: keep the name
        app = SheetApp(self.exam)
        self.assertEqual(app._rename_new_arrivals(["IMG_0001.jpg", "IMG_0002.jpg"]), ["IMG_0001_1.jpg", "IMG_0002.jpg"])
        self.assertEqual((self.success / "IMG_0001.jpg").read_bytes(), first)
        self.assertEqual((self.exam / "IMG_0001_1.jpg").read_bytes(), second)

    def test_regrade_keeps_the_report_and_the_student_bound_in_review(self):
        app, scans = self.grade({"IMG_0001.jpg": ocr("张三")})
        report = self.reports / "IMG_0001.json"
        app.grader_engine = FakeGrader({"IMG_0001.jpg": ocr("张三", exam_id="12345678")})
        # Chosen in review, and an older report from before match records
        for match in ({'method': 'manual', 'candidates': []}, None):
            with self.subTest(match=match):
                data = json.loads(report.read_text(encoding='utf-8'))
                data['db_student_info'] = {'id': '', 'name': '张三', 'class': '701', 'room': '', 'seat': ''}
                data.pop('student_match', None)
                if match:
                    data['student_match'] = match
                report.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
                self.assertEqual(app.regrade_single_file(str(self.success / "IMG_0001.jpg")), (True, "Success"))
                self.assertEqual(self.report_owners(), {"IMG_0001.json": ("IMG_0001.jpg", None)})
                saved = json.loads(report.read_text(encoding='utf-8'))
                self.assertEqual((saved['db_student_info']['name'], saved['student_match']['method']), ('张三', 'manual'))

    def test_review_binds_the_chosen_class_and_saves_in_place(self):
        roster = [{'id': '', 'name': '张三', 'class': '701', 'room': '', 'seat': ''},
                  {'id': '', 'name': '张三', 'class': '702', 'room': '', 'seat': ''},
                  {'id': '', 'name': '王五', 'class': '702', 'room': '03', 'seat': '12'},
                  {'id': '', 'name': '李四', 'class': '701', 'room': '01', 'seat': '6'}]
        sheets = {"IMG_0001.jpg": ocr("张三", "702"), "03-12.jpg": ocr("王五", "702"), "1-6.jpg": ocr("张三", "701")}
        app, scans = self.grade(sheets, students=roster)
        self.assertEqual(set(self.report_owners()), {"IMG_0001.json", "03-12.json", "01-6.json"})
        seat_report = (self.reports / "03-12.json").read_text(encoding='utf-8')

        review = ReviewStub(app, self.reports / "IMG_0001.json")
        review.bind_selected_student("张三 | 702班 |  | -")
        self.assertEqual(review.current_data['db_student_info']['class'], '702')
        # Binding a student whose own Room-Seat report exists must not touch that report
        review.bind_selected_student(ReviewWindow.student_display(roster[2]))
        review.save_to_disk()
        app.on_review_save(review.current_data)

        # Rebinding a Room-Seat sheet to a student without a seat must not leave its report behind
        review = ReviewStub(app, self.reports / "01-6.json")
        review.bind_selected_student(ReviewWindow.student_display(roster[0]))
        review.save_to_disk()
        app.on_review_save(review.current_data)

        self.assertEqual(sorted(os.listdir(self.reports)),
                         ["01-6.json", "01-6.md", "03-12.json", "03-12.md", "IMG_0001.json", "IMG_0001.md"])
        self.assertEqual((self.reports / "03-12.json").read_text(encoding='utf-8'), seat_report)
        saved = json.loads((self.reports / "IMG_0001.json").read_text(encoding='utf-8'))
        self.assertEqual((saved['db_student_info']['name'], saved['review_count']), ('王五', 1))
        saved = json.loads((self.reports / "01-6.json").read_text(encoding='utf-8'))
        self.assertEqual((saved['db_student_info']['name'], saved['db_student_info']['class']), ('张三', '701'))
        self.assertEqual(app.student_manager.students, roster)

    def test_report_lookup_never_returns_another_sheets_report(self):
        self.reports.mkdir()
        (self.reports / "01-6.json").write_text(json.dumps({"original_filename": "1-6.jpg"}), encoding='utf-8')
        (self.reports / "05-01.json").write_text(json.dumps({"total_score": 3}), encoding='utf-8')  # Older report
        reports = str(self.reports)
        self.assertEqual(find_report(reports, "1-6.jpg", ["01-6", "1-6"]), str(self.reports / "01-6.json"))
        self.assertIsNone(find_report(reports, "01-06.jpg", ["01-6", "01-06"]))
        self.assertEqual(find_report(reports, "5-1.jpg", ["05-01", "5-1"]), str(self.reports / "05-01.json"))
        self.assertEqual(unique_report_prefix(reports, "01-6", "1-6.jpg"), "01-6")
        self.assertEqual(unique_report_prefix(reports, "01-6", "01-06.jpg"), "01-06")


class TestRosterFile(unittest.TestCase):
    def test_blank_cells_and_blank_lines(self):
        with tempfile.TemporaryDirectory() as directory:
            xlsx = Path(directory) / "roster.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["考号", "班级", "姓名", "考场", "座号"])
            sheet.append([23090828.0, 701, "张三", None, None])
            sheet.append([None, None, None, None, None])
            sheet.append([None, 702, "张三", None, None])
            workbook.save(xlsx)
            manager = StudentManager()
            self.assertEqual(manager.load_from_file(str(xlsx)), 2)
            self.assertEqual(manager.students, [
                {'id': '23090828', 'name': '张三', 'class': '701', 'room': '', 'seat': ''},
                {'id': '', 'name': '张三', 'class': '702', 'room': '', 'seat': ''},
            ])

            roster_csv = Path(directory) / "roster.csv"
            roster_csv.write_text("班级,姓名\n701,张三\n\n701,李四\n", encoding="utf-8")
            self.assertEqual(manager.load_from_file(str(roster_csv)), 2)
            self.assertEqual([s['name'] for s in manager.students], ['张三', '李四'])

    def load(self, sheets, active=0, name="roster.xlsx"):
        """Load a workbook given as {sheet title: rows}."""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / name
            workbook = openpyxl.Workbook()
            workbook.remove(workbook.active)
            for title, rows in sheets.items():
                sheet = workbook.create_sheet(title)
                for row in rows:
                    sheet.append(row)
            workbook.active = active
            workbook.save(path)
            manager = StudentManager()
            manager.load_from_file(str(path))
            return [(s['class'], s['name'], s['id']) for s in manager.students]

    def test_header_below_title_rows_and_class_from_title(self):
        # Like the rosters exported per class: title, a note, then the header on row 3
        rows = [["高二1班 名单"], ["请核对学生姓名是否与系统一致", "说明：学号以教务系统导出的为准"],
                ["序号", "姓  名", "性别", "成绩", "学号", "状态", "系统编号（勿改）"],
                [1, "张三", "男", None, "2023001", "在班", "x1"], [2, "李四", "女", None, None, "在班", "x2"],
                ["序号", "姓名", "性别", "成绩", "学号", "状态", "系统编号（勿改）"],  # Repeated on a later page
                [3, "王五", "男", None, "2023003", "在班", "x3"], ["合计：3人"]]
        self.assertEqual(self.load({"名单": rows}), [('高二1班', '张三', '2023001'), ('高二1班', '李四', ''), ('高二1班', '王五', '2023003')])

    def test_roster_on_a_later_sheet(self):
        seating = [["讲台"], ["张三", "李四", None, "王五"], ["赵六", None, "钱七", "孙八"]]
        roster = [["高三1班学生名单（按座位排序）"], [], ["序号", "姓名", "排次", "列组"], [1, "张三", "第1排", "第1组"], [2, "李四", "第1排", "第2组"]]
        self.assertEqual(self.load({"座次表": seating, "学生名单": roster}), [('高三1班', '张三', ''), ('高三1班', '李四', '')])

    def test_one_sheet_per_class(self):
        sheets = {"高二1班": [["学号", "姓名"], ["01", "张三"]], "高二2班": [["学号", "姓名"], ["02", "李四"]],
                  "教师": [["姓名", "科目"], ["周老师", "地理"]]}  # Laid out differently: not students
        self.assertEqual(self.load(sheets), [('高二1班', '张三', '01'), ('高二2班', '李四', '02')])
        self.assertEqual(self.load({"Sheet1": [["学号", "姓名"], ["01", "张三"]]}, name="名单-七年级3班.xlsx"), [('七年级3班', '张三', '01')])

    def test_missing_header_is_reported(self):
        with self.assertRaisesRegex(Exception, "姓名"):
            self.load({"座次表": [["讲台"], ["张三", "李四"], ["王五", "赵六"]]})


if __name__ == "__main__":
    unittest.main()
